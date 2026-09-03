import unittest
from io import BytesIO

from app import app, session_data
from openpyxl import Workbook


class LinearRegressionLearningApiTests(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True)
        self.client = app.test_client()
        session_data['datasets'].clear()
        session_data['models'].clear()

    def test_learning_namespace_isolated_and_point_results_are_public(self):
        generated = self.client.post('/api/generate_data', json={
            'type': 'regression',
            'shape': 'linear',
            'n_samples': 30,
            'noise': 0.12,
            'random_state': 42,
            'session_namespace': 'linear_learning',
        })
        self.assertEqual(generated.status_code, 200)
        session_key = generated.get_json()['session_key']
        self.assertEqual(session_key, 'regression_linear_linear_learning')

        trained = self.client.post('/api/linear_regression/train', json={
            'session_key': session_key,
            'algorithm': 'linear',
            'params': {
                'test_size': 0.25,
                'fit_intercept': True,
                'optimizer': 'ols',
            },
        })
        self.assertEqual(trained.status_code, 200)
        results = trained.get_json()['results']
        self.assertEqual(len(results['point_results']), 30)
        self.assertEqual(
            [point['index'] for point in results['point_results']],
            list(range(30)),
        )
        self.assertEqual(set(results['loss_summary']), {
            'sse', 'mse', 'mean_target', 'tss', 'r2',
        })

    def test_linear_regression_page_exposes_compact_free_experiment_workbench(self):
        response = self.client.get('/html/linear_regression.html')
        self.assertEqual(response.status_code, 200)
        page = response.get_data(as_text=True)
        self.assertIn('class="lr-config-rail"', page)
        self.assertIn('role="tablist" aria-label="数据来源"', page)
        self.assertIn('id="excelDropzone"', page)
        self.assertIn('id="excelConfigPanel"', page)
        self.assertIn('id="resultsWorkspace"', page)

    def test_legacy_session_key_format_is_unchanged(self):
        generated = self.client.post('/api/generate_data', json={
            'type': 'regression',
            'shape': 'linear',
            'n_samples': 20,
            'noise': 0.1,
            'random_state': 42,
        })
        self.assertEqual(generated.get_json()['session_key'], 'regression_linear')

    def make_workbook(self, include_invalid=False):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = '训练数据'
        worksheet.append(['面积', '房龄', '房价'])
        for index in range(12):
            area = 60 + index * 5
            age = 2 + index
            price = 30 + area * 0.8 - age * 0.4
            worksheet.append([area, age, price])
        if include_invalid:
            worksheet.append(['无效', 8, 90])
        second = workbook.create_sheet('备注')
        second.append(['说明'])
        second.append(['仅用于预览'])
        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        return stream

    def test_excel_preview_returns_sheets_and_numeric_columns(self):
        response = self.client.post(
            '/api/linear_regression/excel/preview',
            data={'file': (self.make_workbook(), 'sample.xlsx')},
            content_type='multipart/form-data',
        )
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertEqual([sheet['name'] for sheet in payload['sheets']], ['训练数据', '备注'])
        sheet = payload['sheets'][0]
        self.assertEqual(sheet['columns'], ['面积', '房龄', '房价'])
        self.assertEqual(sheet['numeric_columns'], ['面积', '房龄', '房价'])
        self.assertEqual(sheet['total_rows'], 12)
        self.assertEqual(len(sheet['preview']), 5)

    def test_excel_import_filters_invalid_rows_and_trains_with_column_names(self):
        imported = self.client.post(
            '/api/linear_regression/excel/import',
            data={
                'file': (self.make_workbook(include_invalid=True), 'sample.xlsx'),
                'sheet_name': '训练数据',
                'feature_columns': ['面积', '房龄'],
                'target_column': '房价',
            },
            content_type='multipart/form-data',
        )
        self.assertEqual(imported.status_code, 200)
        payload = imported.get_json()
        self.assertTrue(payload['session_key'].startswith('regression_excel_'))
        self.assertEqual(payload['data']['n_samples'], 12)
        self.assertEqual(payload['data']['n_features'], 2)
        self.assertEqual(payload['data']['import_summary']['dropped_rows'], 1)

        trained = self.client.post('/api/linear_regression/train', json={
            'session_key': payload['session_key'],
            'algorithm': 'linear',
            'params': {'optimizer': 'ols', 'test_size': 0.2},
        })
        self.assertEqual(trained.status_code, 200)
        results = trained.get_json()['results']
        self.assertEqual(results['training_info']['feature_names'], ['面积', '房龄'])
        self.assertIn('面积', results['model_equation'])
        self.assertIn('房龄', results['model_equation'])
        self.assertIsNone(results['prediction_curve'])

    def test_excel_import_rejects_invalid_selection_and_too_few_rows(self):
        same_column = self.client.post(
            '/api/linear_regression/excel/import',
            data={
                'file': (self.make_workbook(), 'sample.xlsx'),
                'sheet_name': '训练数据',
                'feature_columns': ['面积'],
                'target_column': '面积',
            },
            content_type='multipart/form-data',
        )
        self.assertEqual(same_column.status_code, 400)
        self.assertIn('不能同时', same_column.get_json()['message'])

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(['x', 'y'])
        for index in range(5):
            worksheet.append([index, index * 2])
        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        too_small = self.client.post(
            '/api/linear_regression/excel/import',
            data={
                'file': (stream, 'small.xlsx'),
                'sheet_name': 'Sheet',
                'feature_columns': ['x'],
                'target_column': 'y',
            },
            content_type='multipart/form-data',
        )
        self.assertEqual(too_small.status_code, 400)
        self.assertIn('至少需要 10 行', too_small.get_json()['message'])

    def test_excel_endpoints_reject_non_xlsx_files(self):
        response = self.client.post(
            '/api/linear_regression/excel/preview',
            data={'file': (BytesIO(b'x,y\n1,2'), 'sample.csv')},
            content_type='multipart/form-data',
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('.xlsx', response.get_json()['message'])

    def test_excel_preview_rejects_malformed_and_oversized_files(self):
        malformed = self.client.post(
            '/api/linear_regression/excel/preview',
            data={'file': (BytesIO(b'not an excel workbook'), 'broken.xlsx')},
            content_type='multipart/form-data',
        )
        self.assertEqual(malformed.status_code, 400)
        self.assertIn('无法解析', malformed.get_json()['message'])

        oversized = self.client.post(
            '/api/linear_regression/excel/preview',
            data={'file': (BytesIO(b'0' * (5 * 1024 * 1024 + 1)), 'large.xlsx')},
            content_type='multipart/form-data',
        )
        self.assertEqual(oversized.status_code, 400)
        self.assertIn('5 MB', oversized.get_json()['message'])

    def test_excel_preview_rejects_duplicate_headers(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(['x', 'x', 'y'])
        worksheet.append([1, 2, 3])
        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        response = self.client.post(
            '/api/linear_regression/excel/preview',
            data={'file': (stream, 'duplicate.xlsx')},
            content_type='multipart/form-data',
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('重复列名', response.get_json()['message'])


if __name__ == '__main__':
    unittest.main()
